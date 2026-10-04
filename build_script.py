# -*- coding: utf-8 -*-
"""
台本・プロンプター生成スクリプト（台本工房）
------------------------------------------------
使い方:  python build_script.py data.json [出力フォルダ]
入力   :  仕様書「6. 中間データ（JSON）」の形式
出力   :  {タイトル}_台本.docx / {タイトル}_プロンプター.pptx
必要   :  pip install python-docx python-pptx
"""
import json, sys, os, re, copy, unicodedata
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pptx import Presentation
from pptx.util import Inches, Pt as PPt, Emu
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ============ 共通設定 ============
FONT_JA = "游ゴシック"        # 和文（eastAsia）
FONT_LATIN = "Yu Gothic"      # 欧文
ALL_ID = "ALL"                # 全員
INK = "2B2B2B"; SUB = "6B6B6B"; LINE = "D9D6D0"; HEAD = "EFEDE8"
SONG_BG = "FFF3C4"            # 楽曲行（パフォーマンス）＝黄色
CUE_BG = "EEEEEE"             # キッカケ
# 部署（キッカケの宛先）: キー → (番号の頭文字, 表示名, 色)
DEPTS = {
    "PA":  ("PA", "PA・音響", "E2E8EE"),
    "VJ":  ("VJ", "VJ・映像", "EAE5EF"),
    "照明": ("L",  "照明",     "F5EAD3"),
    "舞台": ("ST", "舞台",     "E3EBE1"),
    "全体": ("Q",  "全体",     "EEEEEE"),
}
DEPT_ALIAS = {"PA": "PA", "音響": "PA", "SE": "PA", "マイク": "PA", "VJ": "VJ", "映像": "VJ", "VTR": "VJ", "VIDEO": "VJ",
              "照明": "照明", "LIGHT": "照明", "LIGHTING": "照明", "舞台": "舞台", "ステージ": "舞台", "STAGE": "舞台", "転換": "舞台"}
def dept_of(it):
    return DEPT_ALIAS.get(str(it.get("dept", "")).strip().upper(), DEPT_ALIAS.get(str(it.get("dept", "")).strip(), "全体"))

def hex_ok(h, default="888888"):
    h = (h or default).replace("#", "").strip()
    return h.upper() if re.fullmatch(r"[0-9A-Fa-f]{6}", h) else default

def mix(h, other, ratio):
    """h に other を ratio の割合で混ぜる"""
    a = [int(h[i:i+2], 16) for i in (0, 2, 4)]
    b = [int(other[i:i+2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x*(1-ratio)+y*ratio):02X}" for x, y in zip(a, b))

def to_sec(t):
    if not t: return 0
    p = [int(x) for x in str(t).split(":")]
    return p[0]*60 + p[1] if len(p) == 2 else p[0]*3600 + p[1]*60 + p[2]

def fmt_clock(sec):
    return f"{sec//3600:02d}:{(sec%3600)//60:02d}"

def fmt_dur(sec):
    return f"{sec//60}:{sec%60:02d}"

def vwidth(s):
    """見た目の文字幅（全角=1.0, 半角=0.55）"""
    return sum(1.0 if unicodedata.east_asian_width(c) in "WFA" else 0.55 for c in s)

# ============ データ読み込み ============
def load(path):
    with open(path, encoding="utf-8") as f:
        d = json.load(f)
    d.setdefault("meta", {}); d.setdefault("cast", []); d.setdefault("songs", [])
    d.setdefault("rundown", []); d.setdefault("prompter", [])
    cast = {c["id"]: c for c in d["cast"]}
    for c in d["cast"]:
        c["color"] = hex_ok(c.get("color"))
        c.setdefault("short", c.get("role") or c.get("name"))
    cast[ALL_ID] = {"id": ALL_ID, "name": "全員", "short": "全員", "color": "FFFFFF"}
    songs = {s["id"]: s for s in d["songs"]}
    # 時刻の自動計算（time が無い行は前の行＋尺で埋める）
    cur = to_sec(d["meta"]["start"] + ":00") if d["meta"].get("start") else 0
    for row in d["rundown"]:
        if not row.get("duration"):
            sids = [c.get("song") for c in row.get("content", []) if c.get("type") == "song"]
            if sids and sids[0] in songs:
                row["duration"] = songs[sids[0]].get("duration")
        if row.get("time"):
            cur = to_sec(row["time"] + ":00")
        else:
            row["time"] = fmt_clock(cur)
        cur += to_sec(row.get("duration") or "0:00")
    d["_end"] = fmt_clock(cur)
    # キッカケ番号（部署ごとに通し番号）
    cnt = {}
    for row in d["rundown"]:
        for it in row.get("content", []):
            if it.get("type") == "cue":
                k = dept_of(it); it["_dept"] = k
                cnt[k] = cnt.get(k, 0) + 1
                it["_no"] = f"{DEPTS[k][0]}-{cnt[k]:02d}"
    cols = d["meta"].get("columns", ["照明"])
    d["_cols"] = [DEPT_ALIAS.get(str(c).upper(), DEPT_ALIAS.get(str(c), None)) for c in cols]
    d["_cols"] = [c for c in dict.fromkeys(d["_cols"]) if c]
    return d, cast, songs

def who_list(who):
    return who if isinstance(who, list) else [who]

# ============ Word：日本語設定 ============
def set_run_font(run, size=None, bold=None, color=None, shade=None):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.insert(0, rFonts)
    for k, v in (("w:ascii", FONT_LATIN), ("w:hAnsi", FONT_LATIN), ("w:eastAsia", FONT_JA), ("w:cs", FONT_LATIN)):
        rFonts.set(qn(k), v)
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rPr.append(lang)
    lang.set(qn("w:val"), "ja-JP"); lang.set(qn("w:eastAsia"), "ja-JP")
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if shade:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), shade)
        rPr.append(shd)
    return run

def setup_japanese(doc):
    st = doc.styles["Normal"]
    st.font.size = Pt(10.5)
    rPr = st.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.insert(0, rFonts)
    for k, v in (("w:ascii", FONT_LATIN), ("w:hAnsi", FONT_LATIN), ("w:eastAsia", FONT_JA), ("w:cs", FONT_LATIN)):
        rFonts.set(qn(k), v)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rFonts.get(qn(a)) is not None: del rFonts.attrib[qn(a)]
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rPr.append(lang)
    lang.set(qn("w:val"), "ja-JP"); lang.set(qn("w:eastAsia"), "ja-JP"); lang.set(qn("w:bidi"), "ar-SA")
    # 文書既定（docDefaults）にも日本語を設定
    styles = doc.styles.element
    rd = styles.find(qn("w:docDefaults"))
    if rd is not None:
        rpr = rd.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            lg = rpr.find(qn("w:lang"))
            if lg is None:
                lg = OxmlElement("w:lang"); rpr.append(lg)
            lg.set(qn("w:val"), "ja-JP"); lg.set(qn("w:eastAsia"), "ja-JP")
            rf = rpr.find(qn("w:rFonts"))
            if rf is not None:
                for a in list(rf.attrib): del rf.attrib[a]
                rf.set(qn("w:ascii"), FONT_LATIN); rf.set(qn("w:hAnsi"), FONT_LATIN); rf.set(qn("w:eastAsia"), FONT_JA)
    # 設定：テーマ言語を日本語に
    settings = doc.settings.element
    z = settings.find(qn("w:zoom"))
    if z is not None and z.get(qn("w:percent")) is None:
        z.set(qn("w:percent"), "100")
    tfl = settings.find(qn("w:themeFontLang"))
    if tfl is None:
        tfl = OxmlElement("w:themeFontLang"); settings.append(tfl)
    tfl.set(qn("w:val"), "ja-JP"); tfl.set(qn("w:eastAsia"), "ja-JP")
    for s in doc.styles:
        try:
            if s.type == 1 and s.name != "Normal":
                r = s.element.get_or_add_rPr()
                rf = r.find(qn("w:rFonts"))
                if rf is None:
                    rf = OxmlElement("w:rFonts"); r.insert(0, rf)
                for a in list(rf.attrib): del rf.attrib[a]
                rf.set(qn("w:ascii"), FONT_LATIN); rf.set(qn("w:hAnsi"), FONT_LATIN); rf.set(qn("w:eastAsia"), FONT_JA)
        except Exception:
            pass

def cell_shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")): tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def para_shade(p, fill):
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)

def table_borders(table, color=LINE):
    tblPr = table._element.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)

def set_widths(table, widths_mm):
    table.autofit = False
    tblPr = table._element.tblPr
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout"); tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    grid = table._element.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_mm):
        gc.set(qn("w:w"), str(int(w * 56.7)))
    for row in table.rows:
        for c, w in zip(row.cells, widths_mm):
            c.width = Mm(w)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)

def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit"); c.set(qn("w:val"), "true"); trPr.append(c)

def p_fmt(p, before=0, after=2, line=1.25, indent_mm=None, hang_mm=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    if hang_mm is not None:
        pf.left_indent = Mm(hang_mm); pf.first_line_indent = Mm(-hang_mm)
    elif indent_mm is not None:
        pf.left_indent = Mm(indent_mm)
    return p

def add_text(cell_or_doc, text, size=10.5, bold=False, color=INK, first=False, **kw):
    if first and hasattr(cell_or_doc, "paragraphs") and cell_or_doc.paragraphs and not cell_or_doc.paragraphs[0].text:
        p = cell_or_doc.paragraphs[0]
    else:
        p = cell_or_doc.add_paragraph()
    p_fmt(p, **kw)
    if text:
        set_run_font(p.add_run(text), size, bold, color)
    return p

def page_number_footer(section, title):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(p.add_run(f"{title}　"), 8, False, SUB)
    r = p.add_run(); set_run_font(r, 8, False, SUB)
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if t:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), t); r._element.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._element.append(it)

def speaker_label(cast, who):
    return "・".join(cast.get(w, {"short": w})["short"] for w in who_list(who))


# 要素の並び順を Word の仕様どおりに整える（破損防止）
ORDER = {
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "trPr": "cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader tblCellSpacing jc hidden",
}
def fix_order(root):
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for tag, seq in ORDER.items():
        rank = {n: i for i, n in enumerate(seq.split())}
        for el in root.iter(W + tag):
            kids = list(el)
            kids.sort(key=lambda k: rank.get(k.tag.replace(W, ""), 999))
            for k in kids: el.remove(k)
            for k in kids: el.append(k)

# ============ Word 生成 ============
def build_docx(d, cast, songs, out):
    m = d["meta"]
    doc = Document()
    setup_japanese(doc)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    for a in ("left_margin", "right_margin"): setattr(sec, a, Mm(15))
    sec.top_margin, sec.bottom_margin = Mm(14), Mm(14)
    page_number_footer(sec, m.get("title", ""))
    W = 267  # 本文幅 mm

    # --- 表紙ブロック ---
    add_text(doc, m.get("label", "進行台本"), 9, True, SUB, first=True, after=2)
    add_text(doc, m.get("title", "（公演名）"), 22, True, INK, after=2, line=1.1)
    if m.get("subtitle"):
        add_text(doc, m["subtitle"], 12, False, SUB, after=8)
    info = [("日時", m.get("date", "")), ("会場", m.get("venue", "")),
            ("開場／開演", " ／ ".join(x for x in (m.get("open", ""), m.get("start", "")) if x)),
            ("終演予定", m.get("end") or d["_end"]), ("版", m.get("version", ""))]
    info = [x for x in info if x[1]]
    t = doc.add_table(rows=len(info), cols=2); table_borders(t); set_widths(t, [32, 100])
    for i, (k, v) in enumerate(info):
        cell_shade(t.cell(i, 0), HEAD)
        add_text(t.cell(i, 0), k, 9.5, True, SUB, first=True, after=0)
        add_text(t.cell(i, 1), v, 10.5, False, INK, first=True, after=0)

    if m.get("purpose"):
        add_text(doc, "趣旨", 12, True, INK, before=12, after=3)
        for line in str(m["purpose"]).split("\n"):
            add_text(doc, line, 10.5, False, INK, after=2, line=1.4)
    if m.get("notes"):
        add_text(doc, "備考", 12, True, INK, before=10, after=3)
        for n in m["notes"]:
            add_text(doc, "・" + n, 10, False, INK, after=1, line=1.35, hang_mm=4)

    # --- キャスト ---
    if d["cast"]:
        add_text(doc, "キャスト", 12, True, INK, before=12, after=4)
        t = doc.add_table(rows=1, cols=4); table_borders(t); set_widths(t, [12, 60, 60, 60])
        for c, h in zip(t.rows[0].cells, ("色", "表記（台本内）", "演者", "所属・備考")):
            cell_shade(c, HEAD); add_text(c, h, 9, True, SUB, first=True, after=0)
        for c in d["cast"]:
            r = t.add_row().cells
            cell_shade(r[0], mix(c["color"], "FFFFFF", 0.35))
            add_text(r[1], c["short"], 10.5, True, INK, first=True, after=0)
            add_text(r[2], c.get("name", "") + (f"（{c['role']}役）" if c.get("role") else ""), 10, False, INK, first=True, after=0)
            add_text(r[3], c.get("unit", ""), 10, False, SUB, first=True, after=0)

    # --- セットリスト ---
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_text(doc, "セットリスト・タイムテーブル", 14, True, INK, after=6)
    t = doc.add_table(rows=1, cols=5); table_borders(t); set_widths(t, [22, 22, 70, 18, 135])
    for c, h in zip(t.rows[0].cells, ("開始", "区分", "タイトル", "尺", "歌唱・内容")):
        cell_shade(c, HEAD); add_text(c, h, 9, True, SUB, first=True, after=0)
    repeat_header(t.rows[0])
    for row in d["rundown"]:
        song = next((songs.get(x.get("song")) for x in row.get("content", []) if x.get("type") == "song"), None)
        r = t.add_row(); no_split(r); cs = r.cells
        if song: [cell_shade(c, SONG_BG) for c in cs]
        add_text(cs[0], row["time"], 10, True, INK, first=True, after=0)
        add_text(cs[1], row.get("label", ""), 10, True, INK, first=True, after=0)
        add_text(cs[2], f"{song['title']}" if song else row.get("sub", ""), 10, bool(song), INK, first=True, after=0)
        add_text(cs[3], row.get("duration", "") or "", 10, False, SUB, first=True, after=0)
        if song:
            p = cs[4].paragraphs[0]; p_fmt(p, after=0)
            for i, w in enumerate(song.get("singers", [])):
                if i: set_run_font(p.add_run("　"), 10)
                cc = cast.get(w, {"short": w, "color": "CCCCCC"})
                set_run_font(p.add_run(f" {cc['short']} "), 9.5, True, INK, mix(cc["color"], "FFFFFF", 0.6))
            if song.get("credit"):
                set_run_font(p.add_run("　" + song["credit"]), 8.5, False, SUB)
        else:
            add_text(cs[4], row.get("summary", ""), 9.5, False, SUB, first=True, after=0)
    add_text(doc, f"終演予定　{m.get('end') or d['_end']}", 10, True, INK, before=6)

    # --- 進行台本（時間／セットリスト／内容＋部署欄） ---
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_text(doc, "進行台本", 14, True, INK, after=2)
    legend = doc.add_paragraph(); p_fmt(legend, after=6)
    set_run_font(legend.add_run("キッカケ表記　"), 8.5, True, SUB)
    for k, (pre, name, col) in DEPTS.items():
        set_run_font(legend.add_run(f" {pre} "), 8.5, True, INK, col)
        set_run_font(legend.add_run(f" {name}　"), 8.5, False, SUB)
    dcols = d["_cols"]
    cw = 50 if len(dcols) == 1 else 44
    widths = [20, 38] + [0] + [cw] * len(dcols)
    widths[2] = 267 - sum(widths)
    t = doc.add_table(rows=1, cols=3 + len(dcols)); table_borders(t); set_widths(t, widths)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    heads = ["時間", "セットリスト", "内容"] + [DEPTS[k][1] for k in dcols]
    for c, h in zip(t.rows[0].cells, heads):
        cell_shade(c, HEAD); add_text(c, h, 9, True, SUB, first=True, after=0)
    for i, k in enumerate(dcols):
        cell_shade(t.rows[0].cells[3 + i], DEPTS[k][2])
    repeat_header(t.rows[0])
    for row in d["rundown"]:
        r = t.add_row(); cs = r.cells
        has_song = any(x.get("type") == "song" for x in row.get("content", []))
        if has_song:
            cell_shade(cs[0], SONG_BG); cell_shade(cs[1], SONG_BG)
        add_text(cs[0], row["time"], 11, True, INK, first=True, after=0)
        if row.get("duration"):
            add_text(cs[0], f"（{row['duration']}）", 8.5, False, SUB, after=0)
        add_text(cs[1], row.get("label", ""), 11, True, INK, first=True, after=1)
        if row.get("sub"):
            add_text(cs[1], row["sub"], 9, False, SUB, after=0)
        first = True
        dfirst = {k: True for k in dcols}
        for it in row.get("content", []):
            if it.get("type") == "cue" and it["_dept"] in dcols:
                k = it["_dept"]; j = 3 + dcols.index(k)
                render_cue_col(cs[j], it, dfirst[k]); dfirst[k] = False
            else:
                render_item(cs[2], it, cast, songs, first); first = False
    # --- 付録：部署別キューシート ---
    allcues = [(row, it) for row in d["rundown"] for it in row.get("content", []) if it.get("type") == "cue"]
    for k in DEPTS:
        cues = [(row, it) for row, it in allcues if it["_dept"] == k]
        if not cues: continue
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        hp = doc.add_paragraph(); p_fmt(hp, after=6)
        set_run_font(hp.add_run(f" {DEPTS[k][0]} "), 12, True, INK, DEPTS[k][2])
        set_run_font(hp.add_run(f"　キューシート｜{DEPTS[k][1]}"), 14, True, INK)
        ct = doc.add_table(rows=1, cols=5); table_borders(ct); set_widths(ct, [18, 18, 40, 76, 115])
        for c, h in zip(ct.rows[0].cells, ("No", "時間", "セットリスト", "キッカケ", "指示内容")):
            cell_shade(c, HEAD); add_text(c, h, 9, True, SUB, first=True, after=0)
        repeat_header(ct.rows[0])
        for row, it in cues:
            rr = ct.add_row(); no_split(rr); c = rr.cells
            add_text(c[0], it["_no"], 9.5, True, INK, first=True, after=0)
            add_text(c[1], row["time"], 9.5, False, INK, first=True, after=0)
            add_text(c[2], row.get("label", "") + (f"　{row['sub']}" if row.get("sub") else ""), 9.5, False, INK, first=True, after=0)
            add_text(c[3], it.get("trigger") or "（ブロック頭）", 9.5, bool(it.get("trigger")), INK if it.get("trigger") else SUB, first=True, after=0)
            add_text(c[4], it.get("text", ""), 10, False, INK, first=True, after=0, line=1.35)

    # --- 付録：歌詞パート割 ---
    lyric = [s for s in d["prompter"] if s.get("kind") == "lyric"]
    if lyric:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        add_text(doc, "付録　歌詞パート割（プロンプター原稿）", 14, True, INK, after=6)
        cur = None
        for s in lyric:
            if s.get("section") != cur:
                cur = s.get("section")
                add_text(doc, cur, 11.5, True, INK, before=8, after=3)
            if s.get("part"):
                add_text(doc, s["part"], 8.5, False, SUB, before=3, after=1)
            for ln in s["lines"]:
                p = doc.add_paragraph(); p_fmt(p, after=0, line=1.3, hang_mm=22)
                ws = who_list(ln.get("who", ALL_ID))
                cc = cast.get(ws[0], cast[ALL_ID])
                bg = "EEEEEE" if ws[0] == ALL_ID else mix(cc["color"], "FFFFFF", 0.6)
                set_run_font(p.add_run(f" {speaker_label(cast, ws)} "), 8.5, True, INK, bg)
                set_run_font(p.add_run("\t" + ln["text"]), 10.5, False, INK)
                p.paragraph_format.tab_stops.add_tab_stop(Mm(22))
    fix_order(doc.element.body)
    fix_order(doc.styles.element)
    doc.save(out)

def render_cue_col(cell, it, first):
    p = cell.paragraphs[0] if (first and not cell.paragraphs[0].text) else cell.add_paragraph()
    p_fmt(p, before=0 if first else 4, after=0, line=1.2)
    set_run_font(p.add_run(it.get("_no", "")), 8.5, True, SUB)
    if it.get("trigger"):
        q = cell.add_paragraph(); p_fmt(q, after=0, line=1.2)
        set_run_font(q.add_run(it["trigger"]), 8.5, False, SUB)
    q = cell.add_paragraph(); p_fmt(q, after=0, line=1.3)
    set_run_font(q.add_run(it.get("text", "")), 9.5, True, INK)

def render_item(cell, it, cast, songs, first):
    ty = it.get("type")
    def newp():
        if first and not cell.paragraphs[0].text:
            return cell.paragraphs[0]
        return cell.add_paragraph()
    p = newp()
    if ty == "cue":
        k = it.get("_dept", "全体")
        p_fmt(p, before=2, after=2, line=1.25, hang_mm=14); para_shade(p, DEPTS[k][2])
        p.paragraph_format.tab_stops.add_tab_stop(Mm(14))
        set_run_font(p.add_run(it.get("_no", "Q")), 8.5, True, INK)
        set_run_font(p.add_run("\t"), 9)
        if it.get("trigger"):
            set_run_font(p.add_run(it["trigger"] + "　→　"), 9.5, False, SUB)
        set_run_font(p.add_run(it["text"]), 10, True, INK)
    elif ty == "line":
        ws = who_list(it.get("who"))
        p_fmt(p, after=3, line=1.4, hang_mm=20)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(20))
        for i, w in enumerate(ws):
            cc = cast.get(w, {"short": w, "color": "CCCCCC"})
            bg = "E6E6E6" if w == ALL_ID else mix(cc["color"], "FFFFFF", 0.55)
            if i: set_run_font(p.add_run("・"), 9, True, SUB)
            set_run_font(p.add_run(f" {cc['short']} "), 9.5, True, INK, bg)
        set_run_font(p.add_run("\t" + it["text"]), 10.5, False, INK)
    elif ty == "song":
        s = songs.get(it.get("song"), {"title": it.get("song", "?")})
        p_fmt(p, before=2, after=2, line=1.3); para_shade(p, SONG_BG)
        set_run_font(p.add_run(f"♪ {it.get('song','')}「{s['title']}」"), 11, True, INK)
        if s.get("duration"): set_run_font(p.add_run(f"　{s['duration']}"), 9.5, False, SUB)
        if s.get("singers"):
            set_run_font(p.add_run("　歌唱：" + speaker_label(cast, s["singers"])), 9.5, False, SUB)
        if it.get("text"): set_run_font(p.add_run("　" + it["text"]), 9.5, False, SUB)
    elif ty == "topic":
        p_fmt(p, before=2, after=3, line=1.35, hang_mm=6)
        set_run_font(p.add_run("◆ "), 10, True, "8A7A5A")
        set_run_font(p.add_run(it["text"]), 10.5, True, INK)
    elif ty == "direction":
        p_fmt(p, after=3, line=1.35, indent_mm=20)
        set_run_font(p.add_run(f"（{it['text'].strip('（）()')}）"), 9.5, False, SUB)
    else:  # note
        p_fmt(p, after=2, line=1.3)
        set_run_font(p.add_run("※ " + it["text"]), 9, False, SUB)

# ============ PPTX（プロンプター） ============
SW, SH = 13.333, 7.5
BODY_X, BODY_W, NAME_W = 0.6, 12.13, 2.0
BODY_TOP, BODY_BOTTOM = 1.25, 6.45
MAX_PT, MIN_PT = 66, 36
GRAY, DIM, CUEC = "A0A0A0", "6E6E6E", "E8C25A"

def p_run(run, text, size, color, bold=True):
    run.text = text
    f = run.font; f.size = PPt(size); f.bold = bold; f.color.rgb = PRGB.from_string(color)
    f.name = FONT_LATIN
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "ja-JP"); rPr.set("altLang", "en-US")
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); rPr.append(el)
        el.set("typeface", FONT_JA if tag == "a:ea" else FONT_LATIN)
    return run

def tbox(slide, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE, wrap=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.auto_size = None
    return tb, tf

def bright(col):
    """黒背景で読めるように明るく"""
    return "FFFFFF" if col == "FFFFFF" else mix(col, "FFFFFF", 0.30)

def build_pptx(d, cast, songs, out, warnings):
    """プロンプター：黒背景に文字だけ。1スライド＝テキストボックス1つ。人の判別は色のみ。
    現場で書き換えやすいよう、装飾・名前・ヘッダーは入れない（区分やキッカケはノートへ）"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    blank = prs.slide_layouts[6]
    BOX_X, BOX_Y, BOX_W, BOX_H = 0.6, 0.45, SW - 1.2, SH - 0.9
    BASE = 64                      # 標準の文字サイズ（長い行のあるスライドだけ自動で小さくする）

    def color_of(who):
        ws = who_list(who or ALL_ID)
        if len(ws) != 1: return "FFFFFF"
        return bright(cast.get(ws[0], cast[ALL_ID])["color"])

    def new_slide(notes):
        sl = prs.slides.add_slide(blank)
        bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = PRGB(0, 0, 0)
        if notes: sl.notes_slide.notes_text_frame.text = notes
        tb = sl.shapes.add_textbox(Inches(BOX_X), Inches(BOX_Y), Inches(BOX_W), Inches(BOX_H))
        tb.name = "プロンプター本文"
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        bp = tf._txBody.find(qn("a:bodyPr"))
        for el in list(bp): bp.remove(el)
        bp.append(OxmlElement("a:normAutofit"))   # はみ出したらPowerPointが自動で縮小
        return tf

    # 凡例（色と人の対応）
    people = [c for c in d["cast"]]
    if people:
        tf = new_slide("色の凡例（本番前の確認用。不要なら削除）")
        size = min(BASE, int(BOX_H * 72 / (len(people) * 1.5)))
        for k, c in enumerate(people):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            para.line_spacing = 1.15
            p_run(para.add_run(), c.get("short", ""), size, bright(c["color"]))

    slides = d["prompter"]
    for i, s in enumerate(slides):
        kind = s.get("kind", "mc")
        note = "　".join(x for x in (s.get("section", ""), s.get("part", "")) if x)
        if s.get("cue"): note += f"\nキッカケ：{s['cue']}"
        if s.get("notes"): note += "\n" + s["notes"]
        tf = new_slide(note)
        if kind == "title":
            para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
            size = min(66, int(BOX_W * 72 / max(vwidth(s.get("text", "")), 1) * 0.9))
            p_run(para.add_run(), s.get("text", ""), size, "FFFFFF")
            continue
        lines = s.get("lines", [])
        if not lines: continue
        widest = max(vwidth(ln["text"]) for ln in lines)
        size = min(BASE, int(BOX_W * 72 / widest * 0.95), int(BOX_H * 72 / (len(lines) * 1.4)))
        if size < 36:
            warnings.append(f"スライド{i+1}（{s.get('section','')}）: 文字が多すぎます（{size}pt）。行を減らすか分割してください。")
            size = 36
        for k, ln in enumerate(lines):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            para.line_spacing = 1.15
            p_run(para.add_run(), ln["text"], size, color_of(ln.get("who")))
    prs.save(out)

# ============ 実行 ============
def main():
    src = sys.argv[1]
    outdir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(src))
    d, cast, songs = load(src)
    name = re.sub(r'[\\/:*?"<>|\s]+', "_", d["meta"].get("title", "script"))
    docx_path = os.path.join(outdir, f"{name}_台本.docx")
    pptx_path = os.path.join(outdir, f"{name}_プロンプター.pptx")
    warnings = []
    # チェック：未登録キャスト・長すぎる行
    ids = set(cast)
    for i, s in enumerate(d["prompter"]):
        if len(s.get("lines", [])) > 4:
            warnings.append(f"スライド{i+1}（{s.get('section','')}）: {len(s['lines'])}行あります。4行以内に分けてください。")
        for ln in s.get("lines", []):
            for w in who_list(ln.get("who") or ALL_ID):
                if w not in ids: warnings.append(f"未登録のキャストID: {w}")
            if vwidth(ln["text"]) > 18:
                warnings.append(f"長い行（{vwidth(ln['text']):.0f}字幅）: {ln['text']}")
    build_docx(d, cast, songs, docx_path)
    build_pptx(d, cast, songs, pptx_path, warnings)
    print("出力:", docx_path); print("出力:", pptx_path)
    print("終演予定:", d["meta"].get("end") or d["_end"])
    for w in dict.fromkeys(warnings): print("⚠", w)

if __name__ == "__main__":
    main()
